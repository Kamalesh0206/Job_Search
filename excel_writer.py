from pathlib import Path
from datetime import datetime
from typing import List, Dict, Any, Optional, Set
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = [
    "Job ID",
    "Profile ID",
    "Role Category",
    "Job Title",
    "Company",
    "Country",
    "City",
    "Work Arrangement",
    "Remote Eligibility",
    "Visa Sponsorship Status",
    "Salary Range",
    "Experience Required",
    "Required Skills",
    "Matched Skills",
    "Missing Skills",
    "Job Description",
    "Job URL",
    "Source",
    "Posted Date",
    "Found Date",
    "Last Seen Date",
    "Relevance Score",
    "Resume Used",
    "Cover Letter Path",
    "Application Status",
    "Application Date",
    "Recruiter Name",
    "Recruiter Contact",
    "Follow-Up Date",
    "Notes",
    "Fingerprint"
]

VALID_STATUSES = [
    "DISCOVERED",
    "SHORTLISTED",
    "READY_FOR_REVIEW",
    "APPROVED",
    "APPLIED",
    "INTERVIEW",
    "REJECTED",
    "WITHDRAWN",
    "CLOSED"
]

STATUS_MAPPING = {
    "new": "DISCOVERED",
    "discovered": "DISCOVERED",
    "review": "READY_FOR_REVIEW",
    "ready_for_review": "READY_FOR_REVIEW",
    "shortlisted": "SHORTLISTED",
    "approved": "APPROVED",
    "applied": "APPLIED",
    "interview": "INTERVIEW",
    "rejected": "REJECTED",
    "skip": "WITHDRAWN",
    "withdrawn": "WITHDRAWN",
    "offer": "INTERVIEW",
    "closed": "CLOSED"
}

def normalize_status(status: Optional[str]) -> str:
    if not status:
        return "DISCOVERED"
    s = str(status).strip()
    return STATUS_MAPPING.get(s.lower(), s.upper())

def _format_worksheet(ws, max_row: int):
    header_fill = PatternFill("solid", fgColor="1E293B")
    white_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=False)

    for cell in ws[1]:
        cell.font = white_font
        cell.fill = header_fill
        cell.alignment = center_align

    widths = {
        "A": 8,   # Job ID
        "B": 16,  # Profile ID
        "C": 18,  # Role Category
        "D": 38,  # Job Title
        "E": 28,  # Company
        "F": 16,  # Country
        "G": 16,  # City
        "H": 18,  # Work Arrangement
        "I": 18,  # Remote Eligibility
        "J": 22,  # Visa Sponsorship Status
        "K": 16,  # Salary Range
        "L": 18,  # Experience Required
        "M": 26,  # Required Skills
        "N": 26,  # Matched Skills
        "O": 26,  # Missing Skills
        "P": 35,  # Job Description
        "Q": 45,  # Job URL
        "R": 16,  # Source
        "S": 14,  # Posted Date
        "T": 14,  # Found Date
        "U": 14,  # Last Seen Date
        "V": 16,  # Relevance Score
        "W": 30,  # Resume Used
        "X": 35,  # Cover Letter Path
        "Y": 20,  # Application Status
        "Z": 16,  # Application Date
        "AA": 20, # Recruiter Name
        "AB": 24, # Recruiter Contact
        "AC": 16, # Follow-Up Date
        "AD": 25, # Notes
        "AE": 18  # Fingerprint
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:AE{max_row}"

    # Status Data Validation
    dv_formula = f'"{",".join(VALID_STATUSES)}"'
    dv = DataValidation(type="list", formula1=dv_formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"Y2:Y{max(max_row + 500, 5000)}")

def _create_empty(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append(HEADERS)
    _format_worksheet(ws, 1)
    wb.save(path)

def migrate_if_needed(path: Any):
    path = Path(path)
    if not path.exists():
        _create_empty(path)
        return

    wb = load_workbook(path)
    ws = wb["Jobs"] if "Jobs" in wb.sheetnames else wb.active
    
    first_row = [cell.value for cell in ws[1] if cell.value is not None]
    
    # Check if already migrated to new 31-column headers
    if len(first_row) >= len(HEADERS) and first_row[0] == "Job ID" and first_row[1] == "Profile ID":
        wb.close()
        return

    # Need migration from legacy 11 columns
    legacy_rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not any(r):
            continue
        legacy_rows.append(r)

    # Rebuild sheet
    wb.remove(ws)
    new_ws = wb.create_sheet("Jobs", 0)
    new_ws.append(HEADERS)

    for i, r in enumerate(legacy_rows, start=1):
        job_id = r[0] if len(r) > 0 else i
        found_date = str(r[1] or "") if len(r) > 1 else ""
        posted_date = str(r[2] or "") if len(r) > 2 else ""
        raw_role = str(r[3] or "") if len(r) > 3 else ""
        title = str(r[4] or "") if len(r) > 4 else ""
        company = str(r[5] or "") if len(r) > 5 else ""
        location = str(r[6] or "") if len(r) > 6 else ""
        source = str(r[7] or "") if len(r) > 7 else ""
        url = str(r[8] or "") if len(r) > 8 else ""
        raw_status = str(r[9] or "") if len(r) > 9 else "New"
        fp = str(r[10] or "") if len(r) > 10 else ""

        # Map role
        title_l = title.lower()
        role_l = raw_role.lower()
        if "test" in role_l or "test" in title_l or "qa" in role_l or "validation" in title_l or "hil" in title_l or "ecu" in title_l:
            profile_id = "TEST_ENGINEER"
            role_category = "TEST_ENGINEER"
        else:
            profile_id = "AI_ENGINEER"
            role_category = "AI_ENGINEER"

        # Split location
        parts = [p.strip() for p in location.split(",") if p.strip()]
        city = parts[0] if len(parts) >= 1 else ""
        country = parts[-1] if len(parts) >= 2 else location

        status = normalize_status(raw_status)

        new_ws.append([
            job_id,
            profile_id,
            role_category,
            title,
            company,
            country,
            city,
            "Remote" if "remote" in location.lower() or source == "EU Remote Jobs" else "Hybrid / Relocation",
            "Worldwide" if "remote" in location.lower() else "Open",
            "UNKNOWN",
            "",
            "7.4y" if "senior" in title_l else "Flexible",
            "", # required skills
            "Python, AI / Embedded" if profile_id == "AI_ENGINEER" else "Automotive, Python, Testing", # matched
            "", # missing
            f"Discovered listing from {source}. Title: {title} at {company}.",
            url,
            source,
            posted_date,
            found_date or datetime.now().strftime("%Y-%m-%d"),
            found_date or datetime.now().strftime("%Y-%m-%d"),
            75, # Default legacy score
            "resume.pdf",
            "",
            status,
            found_date if status == "APPLIED" else "",
            "",
            "",
            "",
            "Migrated from legacy tracker",
            fp
        ])

    _format_worksheet(new_ws, new_ws.max_row)
    wb.save(path)
    wb.close()

def load_existing_fingerprints(path: Any) -> Set[str]:
    path = Path(path)
    migrate_if_needed(path)
    wb = load_workbook(path, read_only=True)
    ws = wb["Jobs"]
    fps = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        # Fingerprint is the last column (index 30)
        if len(row) >= 31 and row[30]:
            fps.add(str(row[30]))
        elif len(row) >= 11 and row[10]:
            fps.add(str(row[10]))
    wb.close()
    return fps

def load_all_jobs(path: Any) -> List[Dict[str, Any]]:
    path = Path(path)
    migrate_if_needed(path)
    wb = load_workbook(path, read_only=True)
    ws = wb["Jobs"]
    jobs = []
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not any(row) or not row[0]:
            continue
        
        # Pad row to match HEADERS length
        padded = list(row) + [""] * (len(HEADERS) - len(row))
        
        jobs.append({
            "id": padded[0],
            "profile_id": str(padded[1] or "AI_ENGINEER"),
            "role_category": str(padded[2] or "AI_ENGINEER"),
            "title": str(padded[3] or ""),
            "company": str(padded[4] or ""),
            "country": str(padded[5] or ""),
            "city": str(padded[6] or ""),
            "work_arrangement": str(padded[7] or "Remote"),
            "remote_eligibility": str(padded[8] or "Worldwide"),
            "visa_sponsorship_status": str(padded[9] or "UNKNOWN"),
            "salary_range": str(padded[10] or ""),
            "experience_required": str(padded[11] or ""),
            "required_skills": str(padded[12] or ""),
            "matched_skills": str(padded[13] or ""),
            "missing_skills": str(padded[14] or ""),
            "description": str(padded[15] or ""),
            "url": str(padded[16] or ""),
            "source": str(padded[17] or ""),
            "posted_date": str(padded[18] or ""),
            "found_date": str(padded[19] or ""),
            "last_seen_date": str(padded[20] or ""),
            "relevance_score": int(padded[21]) if padded[21] and str(padded[21]).isdigit() else 70,
            "resume_used": str(padded[22] or ""),
            "cover_letter_path": str(padded[23] or ""),
            "status": normalize_status(str(padded[24] or "DISCOVERED")),
            "application_date": str(padded[25] or ""),
            "recruiter_name": str(padded[26] or ""),
            "recruiter_contact": str(padded[27] or ""),
            "follow_up_date": str(padded[28] or ""),
            "notes": str(padded[29] or ""),
            "fingerprint": str(padded[30] or "")
        })
        
    wb.close()
    return jobs

def load_jobs_by_status(path: Path, target_statuses: Optional[List[str]] = None) -> List[Dict[str, Any]]:
    all_jobs = load_all_jobs(path)
    if not target_statuses:
        target_statuses = ["DISCOVERED", "READY_FOR_REVIEW", "SHORTLISTED"]
    targets_norm = {normalize_status(s) for s in target_statuses}
    return [j for j in all_jobs if j["status"] in targets_norm]

def append_jobs(path: Any, rows: List[Dict[str, Any]]) -> int:
    path = Path(path)
    migrate_if_needed(path)
    if not rows:
        return 0

    wb = load_workbook(path)
    ws = wb["Jobs"]
    start_id = ws.max_row

    added_count = 0
    for i, item in enumerate(rows, start=start_id):
        ws.append([
            i,
            item.get("profile_id", "AI_ENGINEER"),
            item.get("role_category", "AI_ENGINEER"),
            item.get("title", ""),
            item.get("company", ""),
            item.get("country", ""),
            item.get("city", ""),
            item.get("work_arrangement", "Remote"),
            item.get("remote_eligibility", "Worldwide"),
            item.get("visa_sponsorship_status", "UNKNOWN"),
            item.get("salary_range", ""),
            item.get("experience_required", ""),
            item.get("required_skills", ""),
            item.get("matched_skills", ""),
            item.get("missing_skills", ""),
            item.get("description", ""),
            item.get("url", ""),
            item.get("source", ""),
            item.get("posted_date", ""),
            item.get("found_date", datetime.now().strftime("%Y-%m-%d")),
            item.get("last_seen_date", datetime.now().strftime("%Y-%m-%d")),
            item.get("relevance_score", 0),
            item.get("resume_used", ""),
            item.get("cover_letter_path", ""),
            normalize_status(item.get("status", "DISCOVERED")),
            item.get("application_date", ""),
            item.get("recruiter_name", ""),
            item.get("recruiter_contact", ""),
            item.get("follow_up_date", ""),
            item.get("notes", ""),
            item.get("fingerprint", "")
        ])
        added_count += 1

    _format_worksheet(ws, ws.max_row)
    wb.save(path)
    wb.close()
    return added_count

def update_job_status(path: Any, job_id: Any, new_status: str, notes: Optional[str] = None) -> bool:
    path = Path(path)
    migrate_if_needed(path)
    wb = load_workbook(path)
    ws = wb["Jobs"]
    
    target_id = str(job_id)
    norm_status = normalize_status(new_status)
    updated = False

    for row in ws.iter_rows(min_row=2):
        if row[0].value and str(row[0].value) == target_id:
            # Column 25 (index 24) is Application Status
            row[24].value = norm_status
            if norm_status == "APPLIED" and not row[25].value:
                # Column 26 (index 25) is Application Date
                row[25].value = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            if notes is not None:
                # Column 30 (index 29) is Notes
                row[29].value = notes
            updated = True
            break

    if updated:
        wb.save(path)
    wb.close()
    return updated

def update_job_application(path: Any, job_id: Any, status: str, resume_used: str, cover_letter_path: str, notes: str = "") -> bool:
    path = Path(path)
    migrate_if_needed(path)
    wb = load_workbook(path)
    ws = wb["Jobs"]
    target_id = str(job_id)
    norm_status = normalize_status(status)
    updated = False

    for row in ws.iter_rows(min_row=2):
        if row[0].value and str(row[0].value) == target_id:
            row[22].value = resume_used # Resume Used
            row[23].value = cover_letter_path # Cover Letter Path
            row[24].value = norm_status # Application Status
            row[25].value = datetime.now().strftime("%Y-%m-%d %H:%M:%S") # Application Date
            if notes:
                row[29].value = notes # Notes
            updated = True
            break

    if updated:
        wb.save(path)
    wb.close()
    return updated

def update_job_last_seen(path: Any, fingerprint: str, last_seen_date: str) -> bool:
    path = Path(path)
    migrate_if_needed(path)
    wb = load_workbook(path)
    ws = wb["Jobs"]
    updated = False

    for row in ws.iter_rows(min_row=2):
        # Column 31 (index 30) is Fingerprint
        if len(row) >= 31 and row[30].value and str(row[30].value) == fingerprint:
            row[20].value = last_seen_date # Last Seen Date
            updated = True
            break

    if updated:
        wb.save(path)
    wb.close()
    return updated
