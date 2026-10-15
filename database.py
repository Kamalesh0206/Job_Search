import os
import sqlite3
import json
import logging
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional, Set

from excel_writer import HEADERS, _format_worksheet

ROOT = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("DATA_DIR", ROOT / "data"))
DATA_DIR.mkdir(parents=True, exist_ok=True)

SQLITE_PATH = DATA_DIR / "jobs.db"
EXCEL_PATH = ROOT / "Job_Search_Tracker.xlsx"

logger = logging.getLogger(__name__)

# Extended headers for Excel export including audit fields
EXTENDED_HEADERS = list(HEADERS)
for col in ["Submission Timestamp", "Submission Response", "Application Reference", "Error Details"]:
    if col not in EXTENDED_HEADERS:
        EXTENDED_HEADERS.append(col)

def get_connection():
    db_url = os.environ.get("DATABASE_URL")
    if db_url and db_url.startswith("postgres"):
        import psycopg2
        import psycopg2.extras
        if db_url.startswith("postgres://"):
            db_url = db_url.replace("postgres://", "postgresql://", 1)
        conn = psycopg2.connect(db_url, cursor_factory=psycopg2.extras.RealDictCursor)
        return conn, "postgres"
    else:
        conn = sqlite3.connect(str(SQLITE_PATH), check_same_thread=False)
        conn.row_factory = sqlite3.Row
        return conn, "sqlite"

def init_db():
    conn, engine = get_connection()
    cursor = conn.cursor()

    if engine == "sqlite":
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS jobs (
                job_id TEXT PRIMARY KEY,
                profile_id TEXT NOT NULL,
                role_category TEXT,
                job_title TEXT NOT NULL,
                company TEXT NOT NULL,
                country TEXT,
                city TEXT,
                work_arrangement TEXT,
                remote_eligibility TEXT,
                visa_sponsorship_status TEXT,
                salary_range TEXT,
                experience_required TEXT,
                required_skills TEXT,
                matched_skills TEXT,
                missing_skills TEXT,
                job_description TEXT,
                job_url TEXT,
                source TEXT,
                posted_date TEXT,
                found_date TEXT,
                last_seen_date TEXT,
                relevance_score INTEGER DEFAULT 0,
                resume_used TEXT,
                cover_letter_path TEXT,
                application_status TEXT DEFAULT 'DISCOVERED',
                application_date TEXT,
                recruiter_name TEXT,
                recruiter_contact TEXT,
                follow_up_date TEXT,
                notes TEXT,
                fingerprint TEXT UNIQUE,
                submission_timestamp TEXT,
                submission_response TEXT,
                application_reference TEXT,
                error_details TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        # Migration: add audit columns if missing
        cursor.execute("PRAGMA table_info(jobs)")
        existing_cols = {row[1] for row in cursor.fetchall()}
        for col_name in ["submission_timestamp", "submission_response", "application_reference", "error_details"]:
            if col_name not in existing_cols:
                cursor.execute(f"ALTER TABLE jobs ADD COLUMN {col_name} TEXT")

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS search_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp TEXT NOT NULL,
                profile TEXT NOT NULL,
                status TEXT NOT NULL,
                jobs_discovered INTEGER DEFAULT 0,
                new_added INTEGER DEFAULT 0,
                duplicates INTEGER DEFAULT 0,
                summary_json TEXT
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS application_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                job_id TEXT NOT NULL,
                event_type TEXT NOT NULL,
                timestamp TEXT NOT NULL,
                notes TEXT
            )
        """)
        conn.commit()
    else:
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS jobs (
                job_id TEXT PRIMARY KEY,
                profile_id TEXT NOT NULL,
                role_category TEXT,
                job_title TEXT NOT NULL,
                company TEXT NOT NULL,
                country TEXT,
                city TEXT,
                work_arrangement TEXT,
                remote_eligibility TEXT,
                visa_sponsorship_status TEXT,
                salary_range TEXT,
                experience_required TEXT,
                required_skills TEXT,
                matched_skills TEXT,
                missing_skills TEXT,
                job_description TEXT,
                job_url TEXT,
                source TEXT,
                posted_date TEXT,
                found_date TEXT,
                last_seen_date TEXT,
                relevance_score INTEGER DEFAULT 0,
                resume_used TEXT,
                cover_letter_path TEXT,
                application_status TEXT DEFAULT 'DISCOVERED',
                application_date TEXT,
                recruiter_name TEXT,
                recruiter_contact TEXT,
                follow_up_date TEXT,
                notes TEXT,
                fingerprint TEXT UNIQUE,
                submission_timestamp TEXT,
                submission_response TEXT,
                application_reference TEXT,
                error_details TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS search_logs (
                id SERIAL PRIMARY KEY,
                timestamp TEXT NOT NULL,
                profile TEXT NOT NULL,
                status TEXT NOT NULL,
                jobs_discovered INTEGER DEFAULT 0,
                new_added INTEGER DEFAULT 0,
                duplicates INTEGER DEFAULT 0,
                summary_json TEXT
            );
            CREATE TABLE IF NOT EXISTS application_events (
                id SERIAL PRIMARY KEY,
                job_id TEXT NOT NULL,
                event_type TEXT NOT NULL,
                timestamp TEXT NOT NULL,
                notes TEXT
            );
        """)
        conn.commit()

    # Check if database table is empty; if so, seed from Excel
    cursor.execute("SELECT COUNT(*) FROM jobs")
    count_row = cursor.fetchone()
    count = count_row[0] if isinstance(count_row, tuple) else (count_row["count"] if "count" in count_row else count_row[0])

    if count == 0 and EXCEL_PATH.exists():
        logger.info("Database is empty. Seeding existing records from Job_Search_Tracker.xlsx...")
        _seed_from_excel(conn, engine)

    conn.close()

def _seed_from_excel(conn, engine):
    from excel_writer import load_all_jobs as excel_load_all
    jobs = excel_load_all(EXCEL_PATH)
    if not jobs:
        return

    cursor = conn.cursor()
    placeholders = ", ".join(["?"] * 31) if engine == "sqlite" else ", ".join(["%s"] * 31)
    insert_sql = f"""
        INSERT OR IGNORE INTO jobs (
            job_id, profile_id, role_category, job_title, company, country, city,
            work_arrangement, remote_eligibility, visa_sponsorship_status, salary_range,
            experience_required, required_skills, matched_skills, missing_skills,
            job_description, job_url, source, posted_date, found_date, last_seen_date,
            relevance_score, resume_used, cover_letter_path, application_status,
            application_date, recruiter_name, recruiter_contact, follow_up_date, notes, fingerprint
        ) VALUES ({placeholders})
    """ if engine == "sqlite" else f"""
        INSERT INTO jobs (
            job_id, profile_id, role_category, job_title, company, country, city,
            work_arrangement, remote_eligibility, visa_sponsorship_status, salary_range,
            experience_required, required_skills, matched_skills, missing_skills,
            job_description, job_url, source, posted_date, found_date, last_seen_date,
            relevance_score, resume_used, cover_letter_path, application_status,
            application_date, recruiter_name, recruiter_contact, follow_up_date, notes, fingerprint
        ) VALUES ({placeholders}) ON CONFLICT (job_id) DO NOTHING
    """

    for idx, j in enumerate(jobs, 1):
        job_id = str(j.get("id") or j.get("job_id") or f"JOB-{idx:04d}")
        score_val = 0
        try:
            score_val = int(j.get("relevance_score") or 0)
        except Exception:
            score_val = 0

        vals = (
            job_id,
            str(j.get("profile_id", "AI_ENGINEER")),
            str(j.get("role_category", "")),
            str(j.get("title") or j.get("job_title", "")),
            str(j.get("company", "")),
            str(j.get("country", "")),
            str(j.get("city", "")),
            str(j.get("work_arrangement", "Remote")),
            str(j.get("remote_eligibility", "Worldwide")),
            str(j.get("visa_sponsorship_status", "UNKNOWN")),
            str(j.get("salary_range", "")),
            str(j.get("experience_required", "")),
            str(j.get("required_skills", "")),
            str(j.get("matched_skills", "")),
            str(j.get("missing_skills", "")),
            str(j.get("description") or j.get("job_description", "")),
            str(j.get("url") or j.get("job_url", "")),
            str(j.get("source", "")),
            str(j.get("posted_date", "")),
            str(j.get("found_date", "")),
            str(j.get("last_seen_date", "")),
            score_val,
            str(j.get("resume_used", "")),
            str(j.get("cover_letter_path", "")),
            str(j.get("status") or j.get("application_status", "DISCOVERED")),
            str(j.get("application_date", "")),
            str(j.get("recruiter_name", "")),
            str(j.get("recruiter_contact", "")),
            str(j.get("follow_up_date", "")),
            str(j.get("notes", "")),
            str(j.get("fingerprint", ""))
        )
        cursor.execute(insert_sql, vals)

    conn.commit()
    logger.info(f"Seeded {len(jobs)} jobs from Excel into SQL database successfully.")

def _row_to_dict(r: Any) -> Dict[str, Any]:
    d = dict(r)
    d["id"] = d.get("job_id", "")
    d["title"] = d.get("job_title", "")
    d["description"] = d.get("job_description", "")
    d["url"] = d.get("job_url", "")
    d["status"] = d.get("application_status", "DISCOVERED")
    return d

def get_existing_fingerprints() -> Set[str]:
    conn, engine = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT fingerprint FROM jobs WHERE fingerprint IS NOT NULL AND fingerprint != ''")
    rows = cursor.fetchall()
    fps = set()
    for r in rows:
        fps.add(r[0] if isinstance(r, tuple) else r["fingerprint"])
    conn.close()
    return fps

def get_all_jobs(
    profile_id: Optional[str] = None,
    status: Optional[str] = None,
    country: Optional[str] = None,
    work_arrangement: Optional[str] = None,
    sponsorship: Optional[str] = None,
    min_score: Optional[int] = None,
    search: Optional[str] = None,
    limit: Optional[int] = None
) -> List[Dict[str, Any]]:
    conn, engine = get_connection()
    cursor = conn.cursor()

    conditions = []
    params = []
    param_marker = "?" if engine == "sqlite" else "%s"

    if profile_id and profile_id != "ALL":
        p_up = profile_id.upper()
        conditions.append(f"(UPPER(profile_id) = {param_marker} OR UPPER(role_category) = 'BOTH')")
        params.append(p_up)

    if status and status != "All":
        conditions.append(f"UPPER(application_status) = {param_marker}")
        params.append(status.upper())

    if country and country != "All":
        conditions.append(f"LOWER(country) LIKE {param_marker}")
        params.append(f"%{country.lower()}%")

    if work_arrangement and work_arrangement != "All":
        conditions.append(f"LOWER(work_arrangement) LIKE {param_marker}")
        params.append(f"%{work_arrangement.lower()}%")

    if sponsorship and sponsorship != "All":
        conditions.append(f"UPPER(visa_sponsorship_status) = {param_marker}")
        params.append(sponsorship.upper())

    if min_score is not None and min_score > 0:
        conditions.append(f"relevance_score >= {param_marker}")
        params.append(min_score)

    if search:
        s_pat = f"%{search.lower()}%"
        conditions.append(f"(LOWER(job_title) LIKE {param_marker} OR LOWER(company) LIKE {param_marker} OR LOWER(country) LIKE {param_marker} OR LOWER(matched_skills) LIKE {param_marker})")
        params.extend([s_pat, s_pat, s_pat, s_pat])

    where_clause = " WHERE " + " AND ".join(conditions) if conditions else ""
    query = f"SELECT * FROM jobs{where_clause} ORDER BY found_date DESC, relevance_score DESC"
    if limit:
        query += f" LIMIT {int(limit)}"

    cursor.execute(query, params)
    rows = cursor.fetchall()
    results = [_row_to_dict(r) for r in rows]
    conn.close()
    return results

def get_job_by_id(job_id: Any) -> Optional[Dict[str, Any]]:
    conn, engine = get_connection()
    cursor = conn.cursor()
    param_marker = "?" if engine == "sqlite" else "%s"
    cursor.execute(f"SELECT * FROM jobs WHERE job_id = {param_marker}", (str(job_id),))
    row = cursor.fetchone()
    conn.close()
    if row:
        return _row_to_dict(row)
    return None

def save_jobs_to_db(job_rows: List[Dict[str, Any]]) -> int:
    if not job_rows:
        return 0
    conn, engine = get_connection()
    cursor = conn.cursor()
    inserted = 0

    placeholders = ", ".join(["?"] * 35) if engine == "sqlite" else ", ".join(["%s"] * 35)
    insert_sql = f"""
        INSERT OR IGNORE INTO jobs (
            job_id, profile_id, role_category, job_title, company, country, city,
            work_arrangement, remote_eligibility, visa_sponsorship_status, salary_range,
            experience_required, required_skills, matched_skills, missing_skills,
            job_description, job_url, source, posted_date, found_date, last_seen_date,
            relevance_score, resume_used, cover_letter_path, application_status,
            application_date, recruiter_name, recruiter_contact, follow_up_date, notes, fingerprint,
            submission_timestamp, submission_response, application_reference, error_details
        ) VALUES ({placeholders})
    """ if engine == "sqlite" else f"""
        INSERT INTO jobs (
            job_id, profile_id, role_category, job_title, company, country, city,
            work_arrangement, remote_eligibility, visa_sponsorship_status, salary_range,
            experience_required, required_skills, matched_skills, missing_skills,
            job_description, job_url, source, posted_date, found_date, last_seen_date,
            relevance_score, resume_used, cover_letter_path, application_status,
            application_date, recruiter_name, recruiter_contact, follow_up_date, notes, fingerprint,
            submission_timestamp, submission_response, application_reference, error_details
        ) VALUES ({placeholders}) ON CONFLICT (job_id) DO NOTHING
    """

    cursor.execute("SELECT COUNT(*) FROM jobs")
    count_row = cursor.fetchone()
    curr_count = count_row[0] if isinstance(count_row, tuple) else (count_row["count"] if "count" in count_row else count_row[0])

    for j in job_rows:
        curr_count += 1
        job_id = str(j.get("id") or j.get("job_id") or f"JOB-{curr_count:04d}")
        score_val = 0
        try:
            score_val = int(j.get("relevance_score") or j.get("match_score") or 0)
        except Exception:
            score_val = 0

        vals = (
            job_id,
            str(j.get("profile_id", "AI_ENGINEER")),
            str(j.get("role_category", "")),
            str(j.get("title") or j.get("job_title", "")),
            str(j.get("company", "")),
            str(j.get("country", "")),
            str(j.get("city", "")),
            str(j.get("work_arrangement", "Remote")),
            str(j.get("remote_eligibility", "Worldwide")),
            str(j.get("visa_sponsorship_status", "UNKNOWN")),
            str(j.get("salary_range", "")),
            str(j.get("experience_required", "")),
            str(j.get("required_skills", "")),
            str(j.get("matched_skills", "")),
            str(j.get("missing_skills", "")),
            str(j.get("description") or j.get("job_description", "")),
            str(j.get("url") or j.get("job_url", "")),
            str(j.get("source", "")),
            str(j.get("posted_date", "")),
            str(j.get("found_date", "")),
            str(j.get("last_seen_date", "")),
            score_val,
            str(j.get("resume_used", "")),
            str(j.get("cover_letter_path", "")),
            str(j.get("status") or j.get("application_status", "DISCOVERED")),
            str(j.get("application_date", "")),
            str(j.get("recruiter_name", "")),
            str(j.get("recruiter_contact", "")),
            str(j.get("follow_up_date", "")),
            str(j.get("notes", "")),
            str(j.get("fingerprint", "")),
            str(j.get("submission_timestamp", "")),
            str(j.get("submission_response", "")),
            str(j.get("application_reference", "")),
            str(j.get("error_details", ""))
        )
        try:
            cursor.execute(insert_sql, vals)
            inserted += 1
        except Exception as e:
            logger.warning(f"Error saving job to DB: {e}")

    conn.commit()
    conn.close()

    try:
        export_to_excel_file(EXCEL_PATH)
    except Exception as e:
        logger.warning(f"Excel sync warning: {e}")

    return inserted

def update_job_status_db(job_id: str, status: str, notes: Optional[str] = None) -> bool:
    conn, engine = get_connection()
    cursor = conn.cursor()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    param_marker = "?" if engine == "sqlite" else "%s"
    if notes:
        sql = f"UPDATE jobs SET application_status = {param_marker}, notes = {param_marker}, updated_at = {param_marker} WHERE job_id = {param_marker}"
        cursor.execute(sql, (status, notes, now_str, str(job_id)))
    else:
        sql = f"UPDATE jobs SET application_status = {param_marker}, updated_at = {param_marker} WHERE job_id = {param_marker}"
        cursor.execute(sql, (status, now_str, str(job_id)))

    # Record application event in audit table
    try:
        sql_event = f"INSERT INTO application_events (job_id, event_type, timestamp, notes) VALUES ({param_marker}, {param_marker}, {param_marker}, {param_marker})"
        cursor.execute(sql_event, (str(job_id), f"STATUS_{status}", now_str, notes or f"Status changed to {status}"))
    except Exception:
        pass

    rowcount = cursor.rowcount
    conn.commit()
    conn.close()

    try:
        export_to_excel_file(EXCEL_PATH)
    except Exception as e:
        logger.warning(f"Excel sync warning: {e}")

    return rowcount > 0

def update_job_application_db(
    job_id: str,
    status: str = "APPLIED",
    resume_used: Optional[str] = None,
    cover_letter_path: Optional[str] = None,
    notes: Optional[str] = None,
    application_date: Optional[str] = None,
    submission_timestamp: Optional[str] = None,
    submission_response: Optional[str] = None,
    application_reference: Optional[str] = None,
    error_details: Optional[str] = None
) -> bool:
    conn, engine = get_connection()
    cursor = conn.cursor()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    app_date = application_date or datetime.now().strftime("%Y-%m-%d")
    sub_time = submission_timestamp or now_str
    param_marker = "?" if engine == "sqlite" else "%s"

    updates = [
        f"application_status = {param_marker}",
        f"application_date = {param_marker}",
        f"submission_timestamp = {param_marker}",
        f"updated_at = {param_marker}"
    ]
    params = [status, app_date, sub_time, now_str]

    if resume_used:
        updates.append(f"resume_used = {param_marker}")
        params.append(resume_used)
    if cover_letter_path:
        updates.append(f"cover_letter_path = {param_marker}")
        params.append(cover_letter_path)
    if notes:
        updates.append(f"notes = {param_marker}")
        params.append(notes)
    if submission_response:
        updates.append(f"submission_response = {param_marker}")
        params.append(submission_response)
    if application_reference:
        updates.append(f"application_reference = {param_marker}")
        params.append(application_reference)
    if error_details:
        updates.append(f"error_details = {param_marker}")
        params.append(error_details)

    params.append(str(job_id))
    sql = f"UPDATE jobs SET {', '.join(updates)} WHERE job_id = {param_marker}"

    cursor.execute(sql, params)
    rowcount = cursor.rowcount

    # Audit event record
    try:
        sql_event = f"INSERT INTO application_events (job_id, event_type, timestamp, notes) VALUES ({param_marker}, {param_marker}, {param_marker}, {param_marker})"
        cursor.execute(sql_event, (str(job_id), f"APPLICATION_{status}", now_str, submission_response or error_details or notes or status))
    except Exception:
        pass

    conn.commit()
    conn.close()

    try:
        export_to_excel_file(EXCEL_PATH)
    except Exception as e:
        logger.warning(f"Excel sync warning: {e}")

    return rowcount > 0

def update_last_seen_db(fingerprints: List[str], date_str: str):
    if not fingerprints:
        return
    conn, engine = get_connection()
    cursor = conn.cursor()
    param_marker = "?" if engine == "sqlite" else "%s"

    for fp in fingerprints:
        sql = f"UPDATE jobs SET last_seen_date = {param_marker} WHERE fingerprint = {param_marker}"
        cursor.execute(sql, (date_str, fp))
    conn.commit()
    conn.close()

def calculate_stats(all_jobs: List[Dict[str, Any]], profile_filter: Optional[str] = None) -> Dict[str, Any]:
    if profile_filter and profile_filter != "ALL":
        target = profile_filter.upper()
        filtered = [j for j in all_jobs if j.get("profile_id", "").upper() == target or j.get("role_category", "").upper() == "BOTH"]
    else:
        filtered = all_jobs

    total = len(filtered)
    relevant = sum(1 for j in filtered if int(j.get("relevance_score", 0)) >= 60)
    high_match = sum(1 for j in filtered if int(j.get("relevance_score", 0)) > 70)
    confirmed_sponsorship = sum(1 for j in filtered if j.get("visa_sponsorship_status", "").upper() == "CONFIRMED")
    remote_jobs = sum(1 for j in filtered if "remote" in j.get("work_arrangement", "").lower())
    applied = sum(1 for j in filtered if j.get("status", "").upper() == "APPLIED")
    ready_to_apply = sum(1 for j in filtered if j.get("status", "").upper() in ("READY_TO_APPLY", "READY_FOR_REVIEW", "APPROVED"))
    submission_failed = sum(1 for j in filtered if j.get("status", "").upper() == "SUBMISSION_FAILED")
    manual_action_required = sum(1 for j in filtered if j.get("status", "").upper() == "MANUAL_ACTION_REQUIRED")
    below_threshold = sum(1 for j in filtered if j.get("status", "").upper() == "BELOW_THRESHOLD")
    interviews = sum(1 for j in filtered if j.get("status", "").upper() == "INTERVIEW")
    rejections = sum(1 for j in filtered if j.get("status", "").upper() == "REJECTED")

    return {
        "total_jobs": total,
        "relevant_jobs": relevant,
        "high_match_jobs": high_match,
        "confirmed_sponsorship": confirmed_sponsorship,
        "remote_jobs": remote_jobs,
        "applications_submitted": applied,
        "ready_to_apply": ready_to_apply,
        "submission_failed": submission_failed,
        "manual_action_required": manual_action_required,
        "below_threshold": below_threshold,
        "interviews": interviews,
        "rejections": rejections,
        "ready_for_review": ready_to_apply,
        "approved": ready_to_apply
    }

def get_stats_db(profile_filter: Optional[str] = "ALL") -> Dict[str, Any]:
    all_jobs = get_all_jobs()
    ai_stats = calculate_stats(all_jobs, "AI_ENGINEER")
    test_stats = calculate_stats(all_jobs, "TEST_ENGINEER")
    overall_stats = calculate_stats(all_jobs, "ALL")

    active_stats = overall_stats
    if profile_filter == "AI_ENGINEER":
        active_stats = ai_stats
    elif profile_filter == "TEST_ENGINEER":
        active_stats = test_stats

    countries = sorted(list({j.get("country", "").strip() for j in all_jobs if j.get("country", "").strip()}))

    return {
        "profile_selected": profile_filter,
        "stats": active_stats,
        "ai_engineer": ai_stats,
        "test_engineer": test_stats,
        "all_profiles": overall_stats,
        "available_countries": countries
    }

def record_search_log(profile: str, status: str, discovered: int, new_added: int, duplicates: int, summary: Dict[str, Any]):
    try:
        conn, engine = get_connection()
        cursor = conn.cursor()
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        param_marker = "?" if engine == "sqlite" else "%s"

        sql = f"""
            INSERT INTO search_logs (timestamp, profile, status, jobs_discovered, new_added, duplicates, summary_json)
            VALUES ({param_marker}, {param_marker}, {param_marker}, {param_marker}, {param_marker}, {param_marker}, {param_marker})
        """
        cursor.execute(sql, (now_str, profile, status, discovered, new_added, duplicates, json.dumps(summary, default=str)))
        conn.commit()
        conn.close()
    except Exception as e:
        logger.warning(f"Error logging search run to DB: {e}")

def get_search_logs(limit: int = 50) -> List[Dict[str, Any]]:
    conn, engine = get_connection()
    cursor = conn.cursor()
    cursor.execute(f"SELECT * FROM search_logs ORDER BY id DESC LIMIT {int(limit)}")
    rows = cursor.fetchall()
    conn.close()
    return [dict(r) for r in rows]

def export_to_excel_file(output_path: Path):
    from openpyxl import Workbook
    conn, engine = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM jobs ORDER BY found_date DESC, relevance_score DESC")
    rows = cursor.fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append(EXTENDED_HEADERS)

    for r in rows:
        d = dict(r)
        ws.append([
            d.get("job_id", ""),
            d.get("profile_id", ""),
            d.get("role_category", ""),
            d.get("job_title", ""),
            d.get("company", ""),
            d.get("country", ""),
            d.get("city", ""),
            d.get("work_arrangement", ""),
            d.get("remote_eligibility", ""),
            d.get("visa_sponsorship_status", ""),
            d.get("salary_range", ""),
            d.get("experience_required", ""),
            d.get("required_skills", ""),
            d.get("matched_skills", ""),
            d.get("missing_skills", ""),
            d.get("job_description", ""),
            d.get("job_url", ""),
            d.get("source", ""),
            d.get("posted_date", ""),
            d.get("found_date", ""),
            d.get("last_seen_date", ""),
            d.get("relevance_score", 0),
            d.get("resume_used", ""),
            d.get("cover_letter_path", ""),
            d.get("application_status", ""),
            d.get("application_date", ""),
            d.get("recruiter_name", ""),
            d.get("recruiter_contact", ""),
            d.get("follow_up_date", ""),
            d.get("notes", ""),
            d.get("fingerprint", ""),
            d.get("submission_timestamp", ""),
            d.get("submission_response", ""),
            d.get("application_reference", ""),
            d.get("error_details", "")
        ])

    _format_worksheet(ws, ws.max_row)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
